#!/usr/bin/env python3
"""
screening.json → 산업 스크리닝 비교 워크북(xlsx).

    python build_screening_workbook.py screening.json [--outdir DIR] [--name 파일명]

시트:
  1. 스크리닝_요약  — 순위표 + 종합점수/상승여력 막대차트
  2. 지표_비교      — 지표 행 × 후보 열(그레이엄 PER×PBR ≤ 22.5 판정은 시트 내 수식)
  3. 점수_산식      — 항목별 점수(값) + 합계/정규화/최종점수는 수식(감사 가능)
  4. 참고_주의      — 가정·기본값 목록, 해자(정성) 기입란, 면책 문구

규칙(리포 공통): 빈 값은 0이 아니라 빈 칸, 판정·합계는 셀 참조 수식, 차트는 막대(+꺾은선만 보조축).
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="4472C4")
TITLE = Font(bold=True, size=13)
NOTE = Font(size=9, color="808080")
BLUE = Font(color="0000CC")  # 정성 기입란(사용자 입력)


def _hdr(ws, row: int, values: list, start_col: int = 1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.font, c.fill = HDR, HDR_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def _num(ws, row, col, v, fmt="#,##0.0"):
    c = ws.cell(row=row, column=col)
    if v is not None:
        c.value = v
        c.number_format = fmt
    return c


def build(data: dict, out_path: Path) -> None:
    cands = data["candidates"]
    n = len(cands)
    wb = Workbook()

    # ---------------- 1. 스크리닝_요약 ----------------
    ws = wb.active
    ws.title = "스크리닝_요약"
    ws["A1"] = f"산업 스크리닝: {data.get('industry', '')} ({data.get('generated_at', '')})"
    ws["A1"].font = TITLE
    ws["A2"] = data.get("scoring_note", "")
    ws["A2"].font = NOTE
    cols = ["순위", "회사", "시장", "기준연도", "현재가", "적정주가(기준)", "상승여력%", "판정",
            "신뢰도", "종합점수", "결측 항목", "경고 수"]
    _hdr(ws, 4, cols)
    r0 = 5
    for i, c in enumerate(cands):
        r = r0 + i
        ws.cell(row=r, column=1, value=c["rank"])
        ws.cell(row=r, column=2, value=c["name"])
        ws.cell(row=r, column=3, value=c["market"])
        ws.cell(row=r, column=4, value=c["metrics"].get("basis_year"))
        _num(ws, r, 5, c["price"], "#,##0.00" if c["currency"] == "USD" else "#,##0")
        _num(ws, r, 6, c["fair_value"], "#,##0.00" if c["currency"] == "USD" else "#,##0")
        _num(ws, r, 7, c["upside_pct"], "+0.0;-0.0")
        ws.cell(row=r, column=8, value=c["verdict"])
        ws.cell(row=r, column=9, value=c["confidence"])
        # 종합점수는 점수_산식 시트의 최종점수 셀 참조(감사 추적)
        ws.cell(row=r, column=10, value=f"=점수_산식!{get_column_letter(2 + i)}{SCORE_FINAL_ROW}").number_format = "0.0"
        ws.cell(row=r, column=11, value=", ".join(c["score"]["missing"]) or "-")
        ws.cell(row=r, column=12, value=len(c["warnings"]))
    for col, w in zip("ABCDEFGHIJKL", (6, 22, 6, 9, 12, 14, 11, 12, 8, 10, 26, 8)):
        ws.column_dimensions[col].width = w

    ch = BarChart(); ch.type = "col"; ch.title = "종합점수(정규화·감점 후)"
    ch.add_data(Reference(ws, min_col=10, min_row=4, max_row=r0 + n - 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=r0, max_row=r0 + n - 1))
    ch.height, ch.width = 8, 16
    ws.add_chart(ch, f"A{r0 + n + 2}")
    ch2 = BarChart(); ch2.type = "col"; ch2.title = "상승여력(%) — 적정주가/현재가−1"
    ch2.add_data(Reference(ws, min_col=7, min_row=4, max_row=r0 + n - 1), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=2, min_row=r0, max_row=r0 + n - 1))
    ch2.height, ch2.width = 8, 16
    ws.add_chart(ch2, f"H{r0 + n + 2}")

    # ---------------- 2. 지표_비교 ----------------
    ws2 = wb.create_sheet("지표_비교")
    ws2["A1"] = "지표 비교 (각 후보의 마지막 '실적' 연도 기준 — (E) 추정연도 제외)"
    ws2["A1"].font = TITLE
    _hdr(ws2, 3, ["지표"] + [c["name"] for c in cands])
    rows = [
        ("ROE 3년 평균(%)", "roe3_pct", "0.0"),
        ("부채비율(%)", "debt_ratio_pct", "0.0"),
        ("이익의 질 OCF/NI(3년)", "ocf_ni3", "0.00"),
        ("FCF 양수 연도 비중", "fcf_pos_share", "0%"),
        ("매출 3년 CAGR(%)", "rev_cagr3_pct", "0.0"),
        ("적자 연도 수(최근≤5년)", "loss_years", "0"),
        ("PER(배)", "per", "0.0"),
        ("PBR(배)", "pbr", "0.00"),
        ("배당수익률(%)", "dividend_yield_pct", "0.00"),
    ]
    r = 4
    per_row = pbr_row = None
    for label, key, fmt in rows:
        ws2.cell(row=r, column=1, value=label).font = Font(bold=True)
        if key == "per":
            per_row = r
        if key == "pbr":
            pbr_row = r
        for i, c in enumerate(cands):
            _num(ws2, r, 2 + i, c["metrics"].get(key), fmt)
        r += 1
    # 그레이엄 PER×PBR — 시트 내 수식으로 계산·판정
    ws2.cell(row=r, column=1, value="그레이엄 PER×PBR").font = Font(bold=True)
    for i in range(n):
        col = get_column_letter(2 + i)
        ws2.cell(row=r, column=2 + i,
                 value=f'=IF(OR({col}{per_row}="",{col}{pbr_row}=""),"",{col}{per_row}*{col}{pbr_row})').number_format = "0.0"
    r += 1
    ws2.cell(row=r, column=1, value="그레이엄 기준(≤22.5)").font = Font(bold=True)
    for i in range(n):
        col = get_column_letter(2 + i)
        ws2.cell(row=r, column=2 + i,
                 value=f'=IF({col}{r-1}="","-",IF({col}{r-1}<=22.5,"충족","미충족"))')
    ws2.cell(row=r + 2, column=1, value="빈 칸 = 해당 데이터 없음(0 아님). PER는 순이익≤0이면 계산하지 않음.").font = NOTE
    ws2.column_dimensions["A"].width = 26
    for i in range(n):
        ws2.column_dimensions[get_column_letter(2 + i)].width = 15

    # ---------------- 3. 점수_산식 ----------------
    ws3 = wb.create_sheet("점수_산식")
    ws3["A1"] = "점수 산식 — 항목 점수(값)는 screening.json에서, 합계·정규화·최종은 수식"
    ws3["A1"].font = TITLE
    _hdr(ws3, 3, ["항목(만점)"] + [c["name"] for c in cands])
    comp_names = [f'{cc["name"]} ({cc["max"]:g})' for cc in cands[0]["score"]["components"]]
    r = 4
    first_comp_row = r
    for j, label in enumerate(comp_names):
        ws3.cell(row=r, column=1, value=label)
        for i, c in enumerate(cands):
            comp = c["score"]["components"][j]
            _num(ws3, r, 2 + i, comp["points"], "0.00")
        r += 1
    last_comp_row = r - 1
    ws3.cell(row=r, column=1, value="가용 만점 합").font = Font(bold=True)
    for i, c in enumerate(cands):
        ws3.cell(row=r, column=2 + i, value=c["score"]["available_max"]).number_format = "0"
    avail_row = r; r += 1
    ws3.cell(row=r, column=1, value="점수 합(수식)").font = Font(bold=True)
    for i in range(n):
        col = get_column_letter(2 + i)
        ws3.cell(row=r, column=2 + i, value=f"=SUM({col}{first_comp_row}:{col}{last_comp_row})").number_format = "0.00"
    sum_row = r; r += 1
    ws3.cell(row=r, column=1, value="정규화 점수(=합/가용만점×100)").font = Font(bold=True)
    for i in range(n):
        col = get_column_letter(2 + i)
        ws3.cell(row=r, column=2 + i, value=f"=IF({col}{avail_row}=0,\"\",{col}{sum_row}/{col}{avail_row}*100)").number_format = "0.0"
    norm_row = r; r += 1
    ws3.cell(row=r, column=1, value="감점(신뢰도)").font = Font(bold=True)
    for i, c in enumerate(cands):
        ws3.cell(row=r, column=2 + i, value=c["score"]["penalty"]).number_format = "0"
    pen_row = r; r += 1
    ws3.cell(row=r, column=1, value="최종점수").font = Font(bold=True)
    for i in range(n):
        col = get_column_letter(2 + i)
        ws3.cell(row=r, column=2 + i, value=f"=IF({col}{norm_row}=\"\",\"\",{col}{norm_row}-{col}{pen_row})").number_format = "0.0"
    assert r == SCORE_FINAL_ROW, f"점수_산식 최종점수 행({r})이 예상({SCORE_FINAL_ROW})과 다릅니다"
    r += 2
    ws3.cell(row=r, column=1, value="항목별 산정 기준").font = Font(bold=True)
    for cc in cands[0]["score"]["components"]:
        r += 1
        ws3.cell(row=r, column=1, value=f'- {cc["name"]}: {cc["note"]}')
    r += 1
    ws3.cell(row=r, column=1, value="- 결측 항목은 0점이 아니라 분모(가용 만점)에서 제외 — 데이터 없음이 점수를 깎지 않되 부풀리지도 않게").font = NOTE
    ws3.column_dimensions["A"].width = 34
    for i in range(n):
        ws3.column_dimensions[get_column_letter(2 + i)].width = 15

    # ---------------- 4. 참고_주의 ----------------
    ws4 = wb.create_sheet("참고_주의")
    ws4["A1"] = "가정·한계·정성 평가"
    ws4["A1"].font = TITLE
    r = 3
    ws4.cell(row=r, column=1, value="경제적 해자(정성) — 점수에 미포함, 최종 추천 시 별도 근거로 기재").font = Font(bold=True)
    r += 1
    _hdr(ws4, r, ["회사", "해자 평가(리서치 후 기입)", "근거/출처"])
    for c in cands:
        r += 1
        ws4.cell(row=r, column=1, value=c["name"])
        ws4.cell(row=r, column=2).font = BLUE
        ws4.cell(row=r, column=3).font = BLUE
    r += 2
    ws4.cell(row=r, column=1, value="후보별 기본값 사용 가정 / 경고").font = Font(bold=True)
    for c in cands:
        r += 1
        ws4.cell(row=r, column=1, value=c["name"]).font = Font(bold=True)
        for t in c["defaults_used"]:
            r += 1
            ws4.cell(row=r, column=1, value=f"  [기본값] {t} — 이 정보들은 정확하지 않습니다")
        for t in c["warnings"]:
            r += 1
            ws4.cell(row=r, column=1, value=f"  [경고] {t}")
    r += 2
    ws4.cell(row=r, column=1, value=data.get("disclaimer", "")).font = NOTE
    ws4.column_dimensions["A"].width = 90
    ws4.column_dimensions["B"].width = 40
    ws4.column_dimensions["C"].width = 40

    wb.save(out_path)
    print(json.dumps({"saved": str(out_path), "candidates": n}, ensure_ascii=False))


SCORE_FINAL_ROW = None  # main에서 계산


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("screening", help="screen_candidates.py가 만든 screening.json")
    ap.add_argument("--outdir", default=".", help="저장 폴더")
    ap.add_argument("--name", help="파일명(기본: 산업명_스크리닝_YYYYMMDD.xlsx)")
    args = ap.parse_args()

    data = json.loads(Path(args.screening).read_text(encoding="utf-8"))
    if not data.get("candidates"):
        raise SystemExit("ERROR: 후보 결과가 비어 있습니다")
    # 점수_산식 시트에서 최종점수가 놓일 행 = 3(헤더) + 항목수 + 4(가용만점·합·정규화·감점) + 1
    n_comp = len(data["candidates"][0]["score"]["components"])
    global SCORE_FINAL_ROW
    SCORE_FINAL_ROW = 3 + n_comp + 4 + 1

    import datetime as dt
    name = args.name or f"{(data.get('industry') or '산업').replace(' ', '')}_스크리닝_{dt.date.today().strftime('%Y%m%d')}.xlsx"
    out = Path(args.outdir) / name
    out.parent.mkdir(parents=True, exist_ok=True)
    build(data, out)


if __name__ == "__main__":
    main()
