"""합성 후보 3사(json 소스)로 스크리닝 점수·순위·워크북 생성을 검증한다."""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

TESTS = Path(__file__).resolve().parent
SCRIPTS = TESTS.parent / "scripts"

BASE = {
    "market": "KR", "currency": "KRW", "amount_unit": 1e8, "unit_label": "억원",
    "years": ["2021", "2022", "2023", "2024", "2025"], "estimated_last": False,
    "price_date": "20260901", "shares": 100_000_000,
}


def _series(rev0, growth, margin, ocf_mult=1.3, capex_ratio=0.4, leverage=0.8):
    rev, ni, eq, li, ocf, cap = [], [], [], [], [], []
    e = rev0 * 0.6
    for i in range(5):
        r = rev0 * (1 + growth) ** i
        n = r * margin
        e = e + n * 0.7
        rev.append(r); ni.append(n); eq.append(e); li.append(e * leverage)
        ocf.append(n * ocf_mult); cap.append(n * ocf_mult * capex_ratio)
    return {"revenue": rev, "ni": ni, "ni_parent": ni, "equity": eq, "equity_parent": eq,
            "assets": [a + b for a, b in zip(eq, li)], "liabilities": li,
            "ocf": ocf, "capex": cap, "da": [r * 0.04 for r in rev],
            "cash": [r * 0.1 for r in rev], "borrowings": [l * 0.5 for l in li],
            "dividends_paid": [max(0.0, n * 0.2) for n in ni]}


def make_fixtures(d: Path) -> Path:
    # A 우량+저평가: 고ROE·저부채·성장, 낮은 주가
    good = dict(BASE, company="A우량", price=9_000, dps=90,
                hist_price=[7000, 7500, 8000, 8500, 9000], series=_series(10000, 0.08, 0.10, leverage=0.6))
    # B 고평가: 비슷한 체력이지만 주가가 매우 높음
    rich = dict(BASE, company="B고평가", price=60_000, dps=90,
                hist_price=[40000, 45000, 50000, 55000, 60000], series=_series(10000, 0.05, 0.08))
    # C 적자·역성장·결측: ni 음수, ocf/capex 결측
    s = _series(10000, -0.05, -0.03, leverage=2.5)
    s["ocf"] = [None] * 5
    s["capex"] = [None] * 5
    bad = dict(BASE, company="C적자", price=5_000, dps=0, hist_price=[None] * 5, series=s)

    cands = []
    for i, fx in enumerate([good, rich, bad]):
        p = d / f"fx{i}.json"
        p.write_text(json.dumps(fx, ensure_ascii=False), encoding="utf-8")
        cands.append({"market": "JSON", "name": fx["company"], "input": str(p)})
    cj = d / "candidates.json"
    cj.write_text(json.dumps({"industry": "테스트산업", "candidates": cands}, ensure_ascii=False), encoding="utf-8")
    return cj


class TestScreener(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        d = Path(cls.tmp.name)
        cj = make_fixtures(d)
        r = subprocess.run([sys.executable, str(SCRIPTS / "screen_candidates.py"), str(cj)],
                           capture_output=True, text=True)
        assert r.returncode == 0, r.stdout + r.stderr
        cls.out = json.loads((d / "screening.json").read_text(encoding="utf-8"))
        r2 = subprocess.run([sys.executable, str(SCRIPTS / "build_screening_workbook.py"),
                             str(d / "screening.json"), "--outdir", str(d), "--name", "t.xlsx"],
                            capture_output=True, text=True)
        assert r2.returncode == 0, r2.stdout + r2.stderr
        cls.xlsx = d / "t.xlsx"

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _by(self, name):
        return next(c for c in self.out["candidates"] if c["name"] == name)

    def test_no_errors_and_all_ranked(self):
        self.assertEqual(self.out["errors"], [])
        self.assertEqual(sorted(c["rank"] for c in self.out["candidates"]), [1, 2, 3])

    def test_ranking_order(self):
        # 우량+저평가 > 고평가 > 적자
        self.assertEqual(self._by("A우량")["rank"], 1)
        self.assertEqual(self._by("C적자")["rank"], 3)
        self.assertGreater(self._by("A우량")["score"]["final"], self._by("B고평가")["score"]["final"])

    def test_missing_excluded_from_denominator(self):
        c = self._by("C적자")
        missing = c["score"]["missing"]
        self.assertIn("이익의 질(OCF/NI)", missing)
        self.assertIn("현금창출(FCF 양수비중)", missing)
        self.assertLess(c["score"]["available_max"], 100)
        # 정규화 점수는 가용 만점 기준
        self.assertAlmostEqual(c["score"]["normalized"],
                               round(c["score"]["raw"] / c["score"]["available_max"] * 100, 1), places=1)

    def test_negative_earnings_no_per(self):
        m = self._by("C적자")["metrics"]
        self.assertIsNone(m["per"])
        self.assertGreaterEqual(m["loss_years"], 1)

    def test_upside_direction(self):
        self.assertGreater(self._by("A우량")["upside_pct"], self._by("B고평가")["upside_pct"])

    def test_workbook_sheets_and_refs(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        self.assertEqual(wb.sheetnames, ["스크리닝_요약", "지표_비교", "점수_산식", "참고_주의"])
        ws = wb["스크리닝_요약"]
        # 종합점수 열은 점수_산식 참조 수식
        self.assertTrue(str(ws.cell(row=5, column=10).value).startswith("=점수_산식!"))
        # 순위 1행 = A우량
        self.assertEqual(ws.cell(row=5, column=2).value, "A우량")
        ws3 = wb["점수_산식"]
        # 최종점수 행에 수식 존재
        found = any(isinstance(c.value, str) and c.value.startswith("=IF(") for row in ws3.iter_rows() for c in row)
        self.assertTrue(found)


if __name__ == "__main__":
    unittest.main()
